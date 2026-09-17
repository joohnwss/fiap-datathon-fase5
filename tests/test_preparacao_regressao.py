import copy
import json
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import dados_pede as d
import auditoria_inicial as a
from preparacao_longitudinal import validate_source_correspondence, write_flat_csv, validate_outputs
from rastreabilidade import write_jsonl, compare_inventories


class PreparacaoRegressionTests(unittest.TestCase):
    def test_observed_phase_codes_and_parentheses(self):
        for label, expected in [("2L", 2), ("3D", 3), ("8E", 8), ("ALFA (2º e 3º ano)", 0),
                                ("Fase 3 (7º e 8º ano)", 3), ("9", 9), (9.0, 9), (" 1a ", 1)]:
            with self.subTest(label=label):
                detail = d.parse_phase_detail(label)
                self.assertEqual(detail["phase"], expected)
                self.assertEqual(detail["original"], label)
        self.assertEqual(d.parse_phase_detail(9)["equivalencia_curricular_status"], "fase_9_sem_significado_documentado")

    def test_phase_rejects_fraction_and_unanchored_digits(self):
        for value in [2.5, "2.5", "FASE 2.5", "2L extra", "série 7", "(Fase 2)", "FASE 2 e 3", "ALFA outra", True, -1, 10, "12A"]:
            with self.subTest(value=value):
                self.assertIsNone(d.parse_phase(value))
        self.assertIsNone(d.safe_parse_int("ano 2022"))
        self.assertIsNone(d.safe_parse_int(2.5))

    def test_cell_states_do_not_fill_unavailable_with_zero(self):
        cases = [(None, "n", "celula_vazia"), (None, "inlineStr", "texto_vazio"),
                 ("", "s", "texto_vazio"), ("  ", "s", "espacos"),
                 ("#N/A", "e", "erro_excel"), ("#DIV/0!", "e", "erro_excel"),
                 ("#N/A", "s", "outro_texto"), ("INCLUIR", "s", "incluir"),
                 (datetime(2023, 1, 1), "d", "data"), (True, "b", "booleano"),
                 ("=1/0", "f", "formula_sem_avaliacao")]
        for value, kind, expected in cases:
            with self.subTest(value=value, kind=kind):
                info = d.classify_cell(value, kind)
                self.assertEqual(info["status"], expected)
                self.assertIsNone(info["numero"])
                self.assertIsNotNone(info["motivo_indisponibilidade"])
        self.assertEqual(d.classify_cell("2,5", "s")["numero"], 2.5)
        self.assertEqual(d.classify_cell(0, "n")["numero"], 0)
        self.assertIsNone(d.classify_cell("1.234,5", "s")["numero"])

    def test_excel_errors_and_types_survive_disk_roundtrip(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "PEDE2024"
        ws.append(["RA", "Fase", "IPP", "IDA", "INDE 2024"])
        ws.append(["A", "2L", "#N/A", "#DIV/0!", "INCLUIR"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "source.xlsx"
            wb.save(path)
            records = d.prepare_records(d.read_frames(path), path)
            r = records[0]
            self.assertIsNone(r["ipp_numerico"])
            self.assertEqual(r["ipp_motivo_indisponibilidade"], "erro_excel:#N/A")
            self.assertEqual(r["originais"]["IPP"]["tipo_excel"], "e")
            self.assertEqual(r["originais"]["IPP"]["valor"], "#N/A")
            self.assertIsNone(r["ida_numerico"])
            self.assertIsNone(r["inde_numerico"])
            validate_source_correspondence(records, path)

    def test_type_after_first_twenty_rows_is_counted(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["IPP"])
        for _ in range(20):
            ws.append([8])
        ws.append(["#N/A"])
        ws.append(["  "])
        frame = d.read_sheet_dataframe(ws)
        types = d.record_type_summary(frame)["IPP"]
        self.assertEqual(types["rows_examined"], 22)
        self.assertEqual(types["type_counts"], {"numero_valido": 20, "erro_excel": 1, "espacos": 1})
        self.assertIn("erro_excel", a.build_field_map({"PEDE2024": frame})[0]["tipos_observados"])

    def test_original_rows_and_all_fields_remain_associated(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "PEDE2024"
        ws.append([None])
        ws.append(["RA", "Fase", "Fase Ideal", "Defasagem", "IAN", "INDE 2024", "INDE 22", "Duplicado", "Duplicado"])
        ws.append(["B", "2L", "Fase 3 (7º ano)", -1, 5, 7, 4, "b1", "b2"])
        ws.append([None])
        ws.append(["A", "3D", "Fase 3 (7º ano)", 0, 10, 8, 3, "a1", "a2"])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "source.xlsx"
            wb.save(path)
            records = d.prepare_records(d.read_frames(path), path)
            self.assertEqual([r["linha_origem"] for r in records], [3, 5])
            self.assertEqual([r["ra"] for r in records], ["B", "A"])
            self.assertEqual(records[1]["inde_numerico"], 8)
            self.assertEqual(records[1]["originais"]["INDE 22"]["valor"], 3)
            self.assertEqual(records[1]["originais"]["Duplicado_1"]["valor"], "a2")
            checks = validate_source_correspondence(records, path)
            self.assertEqual(checks["celulas_originais_conferidas"], 18)
            damaged = copy.deepcopy(records)
            damaged[0]["linha_origem"], damaged[1]["linha_origem"] = 5, 3
            with self.assertRaises(ValueError):
                validate_source_correspondence(damaged, path)
            jsonl, csv_path = Path(tmp) / "result.jsonl", Path(tmp) / "result.csv"
            write_jsonl(jsonl, records)
            write_flat_csv(csv_path, records)
            self.assertTrue(validate_outputs(jsonl, csv_path, records)["jsonl_reaberto_identico"])

    def test_header_suffix_cannot_overwrite_real_column(self):
        names = d.unique_header_names(["RA", "X", "X", "X_1"])
        self.assertEqual(len(names), len(set(names)))
        self.assertEqual(names[-1], "X_1")

    def test_dates_ambiguous_remain_unavailable(self):
        self.assertEqual(d.normalize_birth_date("03/04/2010"), (None, "data_ambigua"))
        self.assertEqual(d.normalize_birth_date("23/04/2010"), ("2010-04-23", "texto_inequivoco"))
        self.assertEqual(d.normalize_birth_date("04/23/2010"), ("2010-04-23", "texto_inequivoco"))
        self.assertEqual(d.normalize_birth_date("31/02/2010"), (None, "data_invalida"))
        self.assertEqual(d.normalize_birth_date("2010-04-23"), ("2010-04-23", "iso"))
        self.assertEqual(d.normalized_registration("genero", " Menina ")[0], "feminino")
        self.assertEqual(d.normalized_registration("instituicao_ensino", "Escola Pública")[0], "pública")

    def test_absent_target_is_not_negative_outcome_and_join_uses_ra(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "PEDE2023"
        ws.append(["RA", "Fase", "Defas"])
        for ra in ("A", "B", "C"):
            ws.append([ra, "ALFA", 0])
        other = wb.create_sheet("PEDE2024")
        other.append(["RA", "Fase", "Defas"])
        other.append(["C", "1A", "#N/A"])
        other.append(["A", "2L", -1])
        records = d.prepare_records({s.title: d.read_sheet_dataframe(s) for s in wb})
        summary, details = a.transition_summary(records, 2023, 2024)
        self.assertEqual(summary["todas_fases"]["with_dest_count"], 2)
        self.assertEqual(summary["todas_fases"]["dest_defasado_count"], 1)
        self.assertEqual(summary["todas_fases"]["found_with_missing_or_invalid_defas_count"], 1)
        for row in details:
            if row["ra"] in {"B", "C"}:
                self.assertIsNone(row["defasagem_futura_binaria"])

    def test_duplicates_kept_but_join_blocked_and_ra_states_separate(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "PEDE2022"
        ws.append(["RA", "Fase", "Defas"])
        for ra in ("A", "A", None, " "):
            ws.append([ra, 1, 0])
        records = d.prepare_records({ws.title: d.read_sheet_dataframe(ws)})
        self.assertEqual(len(records), 4)
        self.assertEqual(sum(r["ra_ano_duplicado"] for r in records), 2)
        self.assertEqual([r["ra_status"] for r in records], ["valido", "valido", "ausente", "vazio_espacos"])
        self.assertTrue(all(r["ipp_status"] == "ausencia_estrutural" for r in records))
        with self.assertRaises(ValueError):
            a.pair_records(records, 2022, 2023)

    def test_all_indicator_ranges_and_extreme_preserved(self):
        for name in d.INDICATORS:
            self.assertEqual(a.operational_column_range(name), (0, 10))
        self.assertIsNone(a.operational_column_range("Destaque IPV"))
        wb = Workbook()
        ws = wb.active
        ws.title = "PEDE2024"
        ws.append(["RA", "Fase", "IAA", "IAN", "Defas"])
        ws.append(["A", "3D", 12.345, None, 0])
        record = d.prepare_records({ws.title: d.read_sheet_dataframe(ws)})[0]
        self.assertEqual(record["iaa_numerico"], 12.345)
        self.assertTrue(record["iaa_fora_faixa_operacional"])
        self.assertIsNone(record["ian_divergente_defasagem"])

    def test_inventory_changes_distinguish_added_removed_changed(self):
        diff = compare_inventories({"a": "1", "b": "2"}, {"b": "3", "c": "4"})
        self.assertEqual(diff, {"adicionados": ["c"], "removidos": ["a"], "alterados": ["b"]})


if __name__ == "__main__":
    unittest.main()
