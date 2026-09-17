"""Leitura e derivações conservadoras compartilhadas por auditoria e preparação.

Nenhuma função salva, recalcula ou modifica o XLSX. Valores brutos e tipos de
célula seguem juntos até a saída JSONL; None representa indisponibilidade.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import unicodedata
from collections import Counter
from datetime import date, datetime
from numbers import Real
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "DATATHON" / "BASE DE DADOS PEDE 2024 - DATATHON.xlsx"
INDICATORS = ("IAN", "IDA", "IEG", "IAA", "IPS", "IPP", "IPV", "INDE")
NUMERIC_PATTERN = re.compile(r"[+-]?(?:\d+(?:[.,]\d+)?|[.,]\d+)(?:[eE][+-]?\d+)?")


def is_missing(value: Any) -> bool:
    return value is None or (not isinstance(value, (list, dict)) and bool(pd.isna(value)))


def json_value(value: Any) -> Any:
    if is_missing(value):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if hasattr(value, "item"):
        return value.item()
    return value


def classify_cell(value: Any, data_type: str | None = None, *, structural: bool = False) -> dict:
    """Distingue conteúdo e tipo original; texto '#N/A' não vira erro Excel.

Números em texto aceitam separador decimal único, sem separador de milhar.
Datas, booleanos e fórmulas não avaliadas nunca viram números implicitamente.
"""
    number = None
    if structural:
        status = "ausencia_estrutural"
    elif data_type == "e":
        status = "erro_excel"
    elif data_type == "f":
        status = "formula_sem_avaliacao"
    elif value is None and data_type == "inlineStr":
        status = "texto_vazio"
    elif is_missing(value):
        status = "celula_vazia"
    elif isinstance(value, bool):
        status = "booleano"
    elif isinstance(value, (datetime, date)) or data_type == "d":
        status = "data"
    elif isinstance(value, Real):
        if math.isfinite(float(value)):
            status, number = "numero_valido", float(value)
        else:
            status = "numero_nao_finito"
    elif isinstance(value, str):
        stripped = value.strip()
        if value == "":
            status = "texto_vazio"
        elif stripped == "":
            status = "espacos"
        elif stripped.upper() == "INCLUIR":
            status = "incluir"
        elif NUMERIC_PATTERN.fullmatch(stripped):
            number = float(stripped.replace(",", "."))
            if math.isfinite(number):
                status = "numero_em_texto"
            else:
                status, number = "numero_nao_finito", None
        else:
            status = "outro_texto"
    else:
        status = "outro_tipo"
    reason = None if number is not None else status
    if status == "erro_excel":
        reason = "erro_excel:" + str(value)
    return {"status": status, "numero": number, "motivo_indisponibilidade": reason}


def safe_parse_int(value: Any) -> int | None:
    number = classify_cell(value)["numero"]
    return int(number) if number is not None and number.is_integer() else None


def parse_phase_detail(value: Any) -> dict:
    phase = None
    status = "nao_interpretavel"
    if is_missing(value):
        status = "ausente"
    elif isinstance(value, Real) and not isinstance(value, bool):
        number = float(value)
        if not math.isfinite(number) or not number.is_integer():
            status = "numerico_nao_inteiro"
        elif 0 <= number <= 9:
            phase, status = int(number), "interpretado_numerico"
        else:
            status = "codigo_fora_dominio_observado"
    elif isinstance(value, str):
        # Apenas anotações finais entre parênteses; nenhum dígito de série entra na fase.
        label = re.sub(r"(?:\s*\([^()]*\))+\s*$", "", value.strip()).strip().upper()
        if not value.strip():
            status = "vazio"
        elif label == "ALFA":
            phase, status = 0, "interpretado_alfa"
        elif re.fullmatch(r"[0-9]", label):
            phase, status = int(label), "interpretado_rotulo"
        elif match := re.fullmatch(r"FASE\s+(?:IDEAL\s+)?([0-9])", label):
            phase, status = int(match[1]), "interpretado_rotulo"
        elif match := re.fullmatch(r"([1-8])([A-Z])", label):
            # Formato observado em PEDE2024: fase de um dígito + letra de turma.
            phase, status = int(match[1]), "interpretado_codigo_turma"
    curricular = "nao_avaliavel" if phase is None else "equivalencia_2022_2024_nao_confirmada"
    if phase == 9:
        curricular = "fase_9_sem_significado_documentado"
    return {"original": value, "phase": phase, "status": status,
            "equivalencia_curricular_status": curricular}


def parse_phase(value: Any) -> int | None:
    return parse_phase_detail(value)["phase"]


def extract_year(sheet_name: str) -> int | None:
    match = re.fullmatch(r"PEDE\s*(20\d{2})", str(sheet_name).strip(), re.IGNORECASE)
    return int(match[1]) if match else None


def classify_ra(value: Any, data_type: str | None = None) -> str:
    if data_type == "e":
        return "erro_excel"
    if value is None and data_type == "inlineStr":
        return "vazio_espacos"
    if is_missing(value):
        return "ausente"
    if isinstance(value, str) and value.strip() == "":
        return "vazio_espacos"
    if isinstance(value, (date, datetime, bool)) or data_type == "f":
        return "tipo_invalido"
    return "valido"


def normalize_ra(value: Any) -> str:
    return "" if is_missing(value) else str(value).strip()


def header_metadata(header: list[Any]) -> list[dict]:
    """Registra repetições antes da renomeação e evita colisões com sufixos reais."""
    bases = [str(v).strip() if v is not None else "" for v in header]
    bases = [v or f"__unnamed_{i}__" for i, v in enumerate(bases)]
    reserved, used, counts, result = set(bases), set(), Counter(), []
    totals = Counter(bases)
    for position, (value, base) in enumerate(zip(header, bases), 1):
        counts[base] += 1
        name = base
        if name in used:
            suffix = counts[base] - 1
            name = f"{base}_{suffix}"
            while name in reserved or name in used:
                suffix += 1
                name = f"{base}_{suffix}"
        used.add(name)
        result.append({"posicao": position, "cabecalho_original": value,
                       "nome_interno": name, "repetido": counts[base] > 1,
                       "ocorrencias_cabecalho": totals[base]})
    return result


def unique_header_names(header: list[Any]) -> list[str]:
    return [m["nome_interno"] for m in header_metadata(header)]


def read_sheet_dataframe(ws) -> pd.DataFrame:
    """Preserva números físicos das linhas; ignora somente linhas sem conteúdo.

Uma linha com espaços ou texto vazio explícito permanece um registro auditável.
O dtype object evita que pandas transforme tipos da fonte antes da auditoria.
"""
    rows = [(n, cells) for n, cells in enumerate(ws.iter_rows(), 1)
            if any(c.value is not None or c.data_type == "inlineStr" for c in cells)]
    if not rows:
        return pd.DataFrame()
    header_row, header_cells = rows[0]
    last_column = max(i for i, c in enumerate(header_cells, 1) if c.value is not None)
    # Colunas além do cabeçalho, caso contenham dados, também são preservadas.
    last_column = max(last_column, max((i for _, row in rows for i, c in enumerate(row, 1)
                                       if c.value is not None), default=0))
    metadata = header_metadata([c.value for c in header_cells[:last_column]])
    data = [[c.value for c in row[:last_column]] for _, row in rows[1:]]
    frame = pd.DataFrame(data, columns=[m["nome_interno"] for m in metadata], dtype=object)
    frame.attrs["header_metadata"] = metadata
    frame.attrs["header_row"] = header_row
    frame.attrs["source_rows"] = [n for n, _ in rows[1:]]
    frame.attrs["physical_max_row"] = ws.max_row
    frame.attrs["cell_metadata"] = {
        m["nome_interno"]: [{"tipo_excel": row[m["posicao"] - 1].data_type,
                            "tipo_python": type(row[m["posicao"] - 1].value).__name__,
                            "formato_numero": row[m["posicao"] - 1].number_format}
                           for _, row in rows[1:]] for m in metadata}
    return frame


def read_frames(path: Path = SOURCE) -> dict[str, pd.DataFrame]:
    wb = load_workbook(path, read_only=False, data_only=False)
    try:
        return {ws.title: read_sheet_dataframe(ws) for ws in wb}
    finally:
        wb.close()


def record_type_summary(frame: pd.DataFrame) -> dict:
    result = {}
    matrix = frame.to_numpy(dtype=object, copy=False)
    for position, col in enumerate(frame.columns):
        metadata = frame.attrs.get("cell_metadata", {}).get(col, [{}] * len(frame))
        counts = Counter(classify_cell(v, m.get("tipo_excel"))["status"]
                         for v, m in zip(matrix[:, position], metadata))
        result[col] = {"observed_types": sorted(counts), "type_counts": dict(counts),
                       "null_count": counts["celula_vazia"], "rows_examined": len(frame)}
    return result


def norm_text(value: Any) -> str | None:
    if is_missing(value) or not str(value).strip():
        return None
    return " ".join(unicodedata.normalize("NFKC", str(value)).casefold().split())


def normalize_birth_date(value: Any) -> tuple[str | None, str]:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d"), "data_excel"
    if is_missing(value) or not str(value).strip():
        return None, "ausente"
    text = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[ T]00:00:00)?", text):
        try:
            return date.fromisoformat(text[:10]).isoformat(), "iso"
        except ValueError:
            return None, "data_invalida"
    if match := re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", text):
        first, second, year = map(int, match.groups())
        candidates = set()
        for month, day in ((second, first), (first, second)):
            try:
                candidates.add(date(year, month, day).isoformat())
            except ValueError:
                pass
        if len(candidates) == 1:
            return candidates.pop(), "texto_inequivoco"
        return None, "data_ambigua" if candidates else "data_invalida"
    return None, "formato_nao_interpretado"


def normalized_registration(field: str, value: Any, data_type: str | None = None) -> tuple[Any, str]:
    if data_type in {"e", "f"}:
        return None, "erro_excel" if data_type == "e" else "formula_sem_avaliacao"
    if field == "data_nascimento":
        return normalize_birth_date(value)
    if field in {"ano_ingresso", "ano_nascimento"}:
        integer = safe_parse_int(value)
        return integer, "inteiro" if integer is not None else "ausente_ou_invalido"
    text = norm_text(value)
    if text is None:
        return None, "ausente"
    if field == "genero":
        genders = {"menina": "feminino", "feminino": "feminino",
                   "menino": "masculino", "masculino": "masculino"}
        return genders.get(text, text), "equivalencia_explicita" if text in genders else "texto_normalizado"
    if field == "instituicao_ensino" and text in {"escola pública", "pública"}:
        return "pública", "equivalencia_explicita_categoria_publica"
    return text, "texto_normalizado"


def column_mapping(frame: pd.DataFrame, year: int) -> dict[str, str | None]:
    lookup = {norm_text(c): c for c in frame.columns}
    aliases = {"ra": ["RA"], "fase": ["Fase"], "fase_ideal": ["Fase ideal"],
               "defasagem": ["Defas", "Defasagem"], "nome": ["Nome", "Nome Anonimizado"],
               "genero": ["Gênero"], "ano_ingresso": ["Ano ingresso"],
               "ano_nascimento": ["Ano nasc"], "data_nascimento": ["Data de Nasc"],
               "idade": ["Idade 22", "Idade"], "instituicao_ensino": ["Instituição de ensino"],
               "escola": ["Escola"], "turma": ["Turma"],
               "pedra": [f"Pedra {year}", f"Pedra {str(year)[2:]}"]}
    for indicator in INDICATORS:
        aliases[indicator.lower()] = [indicator] if indicator != "INDE" else [f"INDE {year}", f"INDE {str(year)[2:]}"]
    return {key: next((lookup[norm_text(a)] for a in candidates if norm_text(a) in lookup), None)
            for key, candidates in aliases.items()}


def expected_ian_value(value):
    d = classify_cell(value)["numero"]
    return None if d is None else 10.0 if d >= 0 else 5.0 if d >= -2 else 2.5


def ian_expected(defasagem: pd.Series) -> pd.Series:
    return pd.Series([expected_ian_value(v) for v in defasagem], index=defasagem.index, dtype="Float64")


def defas_category_counts(series: pd.Series) -> dict:
    d = pd.Series([classify_cell(v)["numero"] for v in series], dtype="Float64")
    return {"D<0": int((d < 0).sum()), "D=0": int((d == 0).sum()),
            "D>0": int((d > 0).sum()), "D>=0": int((d >= 0).sum()),
            "moderada": int(((d >= -2) & (d < 0)).sum()), "severa": int((d < -2).sum()),
            "nao_comparaveis": int(d.isna().sum()), "total": len(d)}


def row_digest(originals: dict) -> str:
    return hashlib.sha256(json.dumps(originals, ensure_ascii=False, sort_keys=True,
                                     allow_nan=False).encode("utf-8")).hexdigest()


def original_record(frame: pd.DataFrame, offset: int, values: dict | None = None) -> dict:
    # iloc em uma linha copiaria recursivamente todos os attrs de todas as células.
    # Acesso escalar/array mantém a mesma associação sem essa cópia quadrática.
    if values is None:
        values = dict(zip(frame.columns, frame.to_numpy(dtype=object, copy=False)[offset]))
    return {m["nome_interno"]: {**m, **frame.attrs["cell_metadata"][m["nome_interno"]][offset],
                                "valor": json_value(values[m["nome_interno"]])}
            for m in frame.attrs["header_metadata"]}


def prepare_records(frames: dict[str, pd.DataFrame], source: Path = SOURCE) -> list[dict]:
    """Uma saída por linha da fonte, sem vincular alunos durante a derivação."""
    try:
        source_name = source.relative_to(ROOT).as_posix()
    except ValueError:
        source_name = str(source)
    records = []
    for sheet, frame in frames.items():
        year = extract_year(sheet)
        if year not in (2022, 2023, 2024):
            raise ValueError(f"Aba fora do escopo da base principal: {sheet}")
        mapping = column_mapping(frame, year)
        matrix = frame.to_numpy(dtype=object, copy=False)
        for offset in range(len(frame)):
            values = dict(zip(frame.columns, matrix[offset]))
            originals = original_record(frame, offset, values)
            record = {"ra": None, "ano_referencia": year, "arquivo_origem": source_name,
                      "aba_origem": sheet, "linha_origem": frame.attrs["source_rows"][offset],
                      "originais": originals, "hash_registro_origem": row_digest(originals)}
            flags = []

            def get(field):
                col = mapping.get(field)
                if col is None:
                    return None, None
                return values[col], originals[col]["tipo_excel"]

            raw_ra, ra_type = get("ra")
            record["ra_status"] = classify_ra(raw_ra, ra_type) if mapping["ra"] else "coluna_ausente"
            if record["ra_status"] == "valido":
                record["ra"] = normalize_ra(raw_ra)
            else:
                flags.append("ra:" + record["ra_status"])
            for field in ("fase", "fase_ideal"):
                value, cell_type = get(field)
                detail = parse_phase_detail(value)
                if cell_type in {"e", "f"}:
                    detail.update(phase=None, status="erro_excel" if cell_type == "e" else "formula_sem_avaliacao",
                                  equivalencia_curricular_status="nao_avaliavel")
                record[field + "_original"] = json_value(value)
                record[field + "_extraida"] = detail["phase"]
                record[field + "_status"] = detail["status"]
                record[field + "_equivalencia_curricular_status"] = detail["equivalencia_curricular_status"]
                if detail["phase"] is None:
                    flags.append(field + ":nao_interpretavel")
                elif detail["phase"] == 9:
                    flags.append(field + ":codigo_9")
            for field in (*[i.lower() for i in INDICATORS], "defasagem", "idade"):
                value, cell_type = get(field)
                detail = classify_cell(value, cell_type,
                                       structural=field == "ipp" and year == 2022 and mapping[field] is None)
                if mapping[field] is None and detail["status"] != "ausencia_estrutural":
                    detail = {"numero": None, "status": "coluna_ausente", "motivo_indisponibilidade": "coluna_ausente"}
                record[field + "_original"] = json_value(value)
                record[field + "_numerico"] = detail["numero"]
                record[field + "_status"] = detail["status"]
                record[field + "_motivo_indisponibilidade"] = detail["motivo_indisponibilidade"]
                if detail["numero"] is None:
                    flags.append(field + ":" + detail["status"])
                if field.upper() in INDICATORS:
                    num = detail["numero"]
                    record[field + "_fora_faixa_operacional"] = None if num is None else not 0 <= num <= 10
                    if record[field + "_fora_faixa_operacional"]:
                        flags.append(field + ":fora_faixa_operacional")
            for field in ("nome", "genero", "ano_ingresso", "ano_nascimento", "data_nascimento",
                          "instituicao_ensino", "escola", "turma", "pedra"):
                value, cell_type = get(field)
                normalized, status = normalized_registration(field, value, cell_type)
                if mapping[field] is None:
                    status = "coluna_ausente"
                record[field + "_padronizado"] = normalized
                record[field + "_status"] = status
                if field == "data_nascimento" and status in {"data_ambigua", "data_invalida", "formato_nao_interpretado"}:
                    flags.append("data_nascimento:" + status)
            # Ano de nascimento pode ser extraído de texto ambíguo apenas quanto ao dia/mês.
            birth_raw, birth_type = get("data_nascimento")
            birth_year = record["ano_nascimento_padronizado"]
            if birth_year is None and birth_type not in {"e", "f"}:
                if isinstance(birth_raw, (date, datetime)):
                    birth_year = birth_raw.year
                elif isinstance(birth_raw, str) and (match := re.fullmatch(r"\d{1,2}[/-]\d{1,2}[/-](\d{4})", birth_raw.strip())):
                    if record["data_nascimento_status"] in {"texto_inequivoco", "data_ambigua"}:
                        birth_year = int(match[1])
                elif record["data_nascimento_padronizado"]:
                    birth_year = int(record["data_nascimento_padronizado"][:4])
                if birth_year is not None:
                    record["ano_nascimento_status"] = "extraido_data_sem_inferir_dia_mes"
            record["ano_nascimento_padronizado"] = birth_year
            d = record["defasagem_numerico"]
            record["defasagem_registrada"] = d
            record["categoria_defasagem"] = (None if d is None else "sem_defasagem" if d >= 0
                                              else "moderada" if d >= -2 else "severa")
            record["ian_esperado"] = expected_ian_value(d)
            ian = record["ian_numerico"]
            record["ian_divergente_defasagem"] = None if d is None or ian is None else ian != record["ian_esperado"]
            phase, ideal = record["fase_extraida"], record["fase_ideal_extraida"]
            calc = None if phase is None or ideal is None else phase - ideal
            record["defasagem_calculada"] = calc
            record["defasagem_calculada_status"] = "aritmetica_de_codigos_sem_confirmacao_curricular" if calc is not None else "nao_comparavel"
            record["defasagem_divergente_calculada"] = None if d is None or calc is None else d != calc
            for flag in ("ian_divergente_defasagem", "defasagem_divergente_calculada"):
                if record[flag]:
                    flags.append(flag)
            record["flags_qualidade"] = flags
            records.append(record)
    counts = Counter((r["ra"], r["ano_referencia"]) for r in records if r["ra"] is not None)
    for record in records:
        record["ra_ano_duplicado"] = record["ra"] is not None and counts[(record["ra"], record["ano_referencia"])] > 1
        if record["ra_ano_duplicado"]:
            record["flags_qualidade"].append("ra_ano_duplicado")
    return records
