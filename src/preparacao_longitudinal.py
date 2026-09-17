"""Base longitudinal conservadora com uma observação por registro anual da fonte."""
from __future__ import annotations
import csv
import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

from dados_pede import (ROOT, SOURCE, INDICATORS, read_frames, prepare_records, json_value, row_digest,
                        header_metadata, column_mapping, extract_year)
import pandas as pd
from auditoria_inicial import audit_frames
from rastreabilidade import start_run, finish_run, write_json, write_jsonl
from relatorios_preparacao import write_field_map, report_body


def validate_source_correspondence(records: list[dict], source: Path = SOURCE) -> dict:
    """Relê o XLSX, confere cada campo e cada endereço, sem reutilizar o leitor tabular."""
    wb = load_workbook(source, read_only=False, data_only=False)
    expected_addresses = set()
    cells_checked = 0
    source_headers = {}
    try:
        for ws in wb:
            content_rows = [n for n, row in enumerate(ws.iter_rows(), 1)
                            if any(c.value is not None or c.data_type == "inlineStr" for c in row)]
            expected_addresses.update((ws.title, n) for n in content_rows[1:])
            if content_rows:
                header_row = content_rows[0]
                headers = header_metadata([c.value for c in ws[header_row]])
                header_frame = pd.DataFrame(columns=[h["nome_interno"] for h in headers])
                ra_name = column_mapping(header_frame, extract_year(ws.title))["ra"]
                ra_position = next((h["posicao"] for h in headers if h["nome_interno"] == ra_name), None)
                source_headers[ws.title] = (header_row, ra_position)
        actual_addresses = [(r["aba_origem"], r["linha_origem"]) for r in records]
        if len(actual_addresses) != len(set(actual_addresses)) or set(actual_addresses) != expected_addresses:
            raise ValueError("Registros preparados não correspondem um a um às linhas físicas da origem")
        for record in records:
            ws = wb[record["aba_origem"]]
            line = record["linha_origem"]
            header_row, ra_position = source_headers[ws.title]
            if record["hash_registro_origem"] != row_digest(record["originais"]):
                raise ValueError("Hash de registro serializado inválido")
            for original in record["originais"].values():
                col = original["posicao"]
                cell = ws.cell(line, col)
                if (original["valor"] != json_value(cell.value)
                    or original["tipo_excel"] != cell.data_type
                    or original["tipo_python"] != type(cell.value).__name__
                    or original["formato_numero"] != cell.number_format
                    or original["cabecalho_original"] != ws.cell(header_row, col).value):
                    raise ValueError(f"Campo divergente da origem: {ws.title}!{cell.coordinate}")
                cells_checked += 1
            if record["ra_status"] == "valido" and (ra_position is None or record["ra"] != str(ws.cell(line, ra_position).value).strip()):
                raise ValueError("RA perdeu associação com a linha de origem")
            for indicator in INDICATORS:
                key = indicator.lower()
                if record[key + "_numerico"] is None:
                    if record[key + "_motivo_indisponibilidade"] is None:
                        raise ValueError("Indicador indisponível sem motivo")
                elif record[key + "_motivo_indisponibilidade"] is not None:
                    raise ValueError("Indicador numérico com motivo indevido de indisponibilidade")
    finally:
        wb.close()
    counts = Counter(r["ano_referencia"] for r in records)
    return {"registros_conferidos_na_origem": len(records), "celulas_originais_conferidas": cells_checked,
            "correspondencia_integral_origem": True, "enderecos_origem_unicos": True,
            "contagens_por_ano": dict(counts),
            "ra_ano_unico": not any(r["ra_ano_duplicado"] for r in records),
            "registros_com_ra_invalido": sum(r["ra_status"] != "valido" for r in records)}


def write_flat_csv(path: Path, records: list[dict]) -> None:
    fields = [f for f in records[0] if f != "originais"]
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for record in records:
            flat = {k: record[k] for k in fields}
            flat["flags_qualidade"] = json.dumps(flat["flags_qualidade"], ensure_ascii=False)
            writer.writerow(flat)


def validate_outputs(jsonl: Path, csv_path: Path, records: list[dict]) -> dict:
    with jsonl.open(encoding="utf-8") as stream:
        reloaded = [json.loads(line) for line in stream]
    if reloaded != records:
        raise ValueError("JSONL não preservou integralmente registros, valores e tipos")
    with csv_path.open(encoding="utf-8", newline="") as stream:
        flat = list(csv.DictReader(stream))
    if len(flat) != len(records):
        raise ValueError("CSV perdeu registros")
    for raw, record in zip(flat, records):
        for key, value in raw.items():
            expected = record[key]
            expected = json.dumps(expected, ensure_ascii=False) if key == "flags_qualidade" else "" if expected is None else str(expected)
            if value != expected:
                raise ValueError(f"CSV alterou o campo {key}")
    return {"jsonl_reaberto_identico": True, "csv_reaberto_campos_conferidos": True,
            "campos_derivados_csv": len(flat[0]), "linhas_jsonl": len(reloaded)}


def main() -> None:
    run = start_run("preparacao")
    frames = read_frames()
    records = prepare_records(frames)
    validations = validate_source_correspondence(records)
    local = ROOT / "local_data"
    jsonl, csv_path = local / "base_longitudinal.jsonl", local / "base_longitudinal.csv"
    write_jsonl(jsonl, records)
    write_flat_csv(csv_path, records)
    validations.update(validate_outputs(jsonl, csv_path, records))
    summary, details = audit_frames(frames, records)
    for kind, rows in details.items():
        write_jsonl(local / "auditoria" / f"detalhes_{kind}.jsonl", rows)
    write_json(local / "auditoria/resumo.json", summary)
    write_field_map(summary["mapa_campos"])
    validations["referencias_conferidas"] = sum(r["confere"] for r in summary["referencias"])
    validations["referencias_total"] = len(summary["referencias"])
    validations["fluxos_transicoes_fecham"] = all(
        t["eligible_count"] == t["with_dest_count"] + t["not_found_in_destination_count"]
        and t["with_dest_count"] == t["found_with_missing_or_invalid_defas_count"] + t["dest_defasado_count"] + t["dest_em_fase_count"] + t["dest_adiantado_count"]
        for group in summary["transicoes"].values() for t in group.values())
    validations["classificacao_indicadores_fecha_por_ano"] = all(sum(q["estados"].values()) == q["total"] for q in summary["indicadores"])
    outputs = [jsonl, csv_path, local / "auditoria/resumo.json", ROOT / "docs/mapa_campos.md"]
    outputs.extend(local / "auditoria" / f"detalhes_{k}.jsonl" for k in details)
    metadata = finish_run(run, summary, validations, outputs)
    report = "# Relatório de preparação inicial\n\n" + report_body(summary, metadata)
    report += "\n## Arquivos produzidos\n\n- Base completa: `local_data/base_longitudinal.jsonl`.\n- Visão plana: `local_data/base_longitudinal.csv`.\n- Detalhes: `local_data/auditoria/detalhes_{celulas,cadastro,transicoes}.jsonl`.\n- Mapa: `docs/mapa_campos.md`.\n- Metadados agregados: `reports/metadados_preparacao.json`.\n"
    (ROOT / "reports/relatorio_preparacao_inicial.md").write_text(report, encoding="utf-8")
    print(f"Preparação: {len(records)} registros; {validations['celulas_originais_conferidas']} células conferidas; referências: {validations['referencias_conferidas']}/{validations['referencias_total']}; integridade: {metadata['source_integrity_preserved']}")


if __name__ == "__main__":
    main()
